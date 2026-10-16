# app/utils/pptx_exporter.py
import os
import re
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from app.core.models import AgentExtractionResult
from app.utils.charts import (
    generate_subsystem_breakdown_chart,
    generate_capability_radar_chart,
    generate_native_blueprint_chart,
    generate_native_sequence_chart
)
from app.utils.diagram_renderer import render_mermaid_to_png

# ── Executive Light Mode Corporate Palette ──
BG_CANVAS = RGBColor(248, 250, 252)     # #F8FAFC (Soft executive slate-50)
CARD_BG = RGBColor(255, 255, 255)       # #FFFFFF (Crisp Pure White)
CARD_BG_ALT = RGBColor(241, 245, 249)   # #F1F5F9 (Slate 100)
TEXT_DARK = RGBColor(15, 23, 42)        # #0F172A (Deep Slate 900)
TEXT_BODY = RGBColor(51, 65, 85)        # #334155 (Slate 700 - high contrast)
TEXT_MUTED = RGBColor(100, 116, 139)    # #64748B (Slate 500)

ACCENT_NAVY = RGBColor(30, 58, 138)     # #1E3A8A (Navy 900)
ACCENT_BLUE = RGBColor(29, 78, 216)     # #1D4ED8 (Royal Blue 700)
ACCENT_CYAN = RGBColor(2, 132, 199)     # #0284C7 (Sky 600)
ACCENT_GREEN = RGBColor(4, 120, 87)     # #047857 (Emerald 700)
ACCENT_PURPLE = RGBColor(109, 40, 217)  # #6D28D9 (Violet 700)
ACCENT_AMBER = RGBColor(180, 83, 9)     # #B45309 (Amber 700)

BORDER_COLOR = RGBColor(203, 213, 225)  # #CBD5E1 (Slate 300)
BORDER_LIGHT = RGBColor(226, 232, 240)  # #E2E8F0 (Slate 200)

FONT_FAMILY = "Segoe UI"


def sanitize_text(text: str) -> str:
    """Masks production secrets and keys from presentation slides."""
    if not text:
        return ""
    text = re.sub(r'(MONGODB_URI=)[^\s,]+', r'\1mongodb+srv://[REDACTED]', text)
    text = re.sub(r'(SMTP_PASS=)[^\s,]+', r'\1[REDACTED_PASSWORD]', text)
    text = re.sub(r'(JWT_SECRET=)[^\s,]+', r'\1[REDACTED_SECRET]', text)
    text = re.sub(r'(DB_PASSWORD=)[^\s,]+', r'\1[REDACTED_PASSWORD]', text)
    return text


def set_slide_background(slide, prs):
    """Sets an executive Light Mode canvas with a top brand accent stripe."""
    # Main background canvas
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_CANVAS
    bg.line.color.rgb = BG_CANVAS

    # Top brand ribbon (Navy/Blue)
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.08))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = ACCENT_BLUE
    top_bar.line.color.rgb = ACCENT_BLUE

    return bg


def add_slide_footer(slide):
    """Adds enterprise footer watermark and standard."""
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.35))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PRISM Architectural Intelligence  •  Tier-1 Production Standard  •  Confidential"
    p.font.size = Pt(9)
    p.font.color.rgb = TEXT_MUTED
    p.font.name = FONT_FAMILY


def add_header(slide, title_text: str, category_text: str = "PRISM ARCHITECTURAL AUDIT"):
    """Adds a modern two-tier header in Light Mode with category chip and bold title."""
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.32), Inches(11.7), Inches(1.05))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p_cat = tf.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10.5)
    p_cat.font.bold = True
    p_cat.font.color.rgb = ACCENT_BLUE
    p_cat.font.name = FONT_FAMILY
    
    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_DARK
    p_title.font.name = FONT_FAMILY
    p_title.space_before = Pt(2)

    add_slide_footer(slide)


def add_card(slide, left, top, width, height, title: str, body_text: str, 
             font_size: float = 12.5, title_color: RGBColor = ACCENT_BLUE,
             bg_color: RGBColor = CARD_BG, border_color: RGBColor = BORDER_COLOR):
    """Draws a clean white container card in Light Mode with large legible text."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    
    # Left accent indicator strip
    accent_strip = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, 
        left + Inches(0.04), top + Inches(0.12), Inches(0.08), height - Inches(0.24)
    )
    accent_strip.fill.solid()
    accent_strip.fill.fore_color.rgb = title_color
    accent_strip.line.color.rgb = title_color

    tb = slide.shapes.add_textbox(
        left + Inches(0.3), top + Inches(0.18), width - Inches(0.55), height - Inches(0.36)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    
    p_title = tf.paragraphs[0]
    p_title.text = title.upper()
    p_title.font.size = Pt(14)
    p_title.font.bold = True
    p_title.font.color.rgb = title_color
    p_title.font.name = FONT_FAMILY
    
    p_body = tf.add_paragraph()
    clean_body = sanitize_text(body_text)
    # Truncate overly long text blocks safely
    if len(clean_body) > 650:
        clean_body = clean_body[:640] + " ... [See Full Report for Details]"
    p_body.text = clean_body
    p_body.font.size = Pt(font_size)
    p_body.font.color.rgb = TEXT_BODY
    p_body.font.name = FONT_FAMILY
    p_body.space_before = Pt(6)


def generate_pptx_deck(result: AgentExtractionResult, owner: str = "", repo: str = "") -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    target_repo = f"{owner}/{repo}" if owner and repo else result.project_name
    deck = result.presentation_deck
    hg = getattr(result, "developer_handover_guide", None)

    # =========================================================
    # 1. Executive Cover Slide (Light Mode)
    # =========================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, prs)
    
    # Hero card framing
    hero_card = slide1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.733), Inches(5.4)
    )
    hero_card.fill.solid()
    hero_card.fill.fore_color.rgb = CARD_BG
    hero_card.line.color.rgb = BORDER_COLOR
    hero_card.line.width = Pt(1.5)

    # Decorative left accent bar on hero card
    hero_accent = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.18), Inches(5.4)
    )
    hero_accent.fill.solid()
    hero_accent.fill.fore_color.rgb = ACCENT_BLUE
    hero_accent.line.color.rgb = ACCENT_BLUE

    t_box = slide1.shapes.add_textbox(Inches(1.4), Inches(1.7), Inches(10.7), Inches(4.4))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p_badge = tf1.paragraphs[0]
    p_badge.text = "ENTERPRISE ARCHITECTURAL AUDIT & ENGINEERING HANDOVER"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = ACCENT_BLUE
    p_badge.font.name = FONT_FAMILY
    
    p_main = tf1.add_paragraph()
    p_main.text = result.project_name
    p_main.font.size = Pt(40)
    p_main.font.bold = True
    p_main.font.color.rgb = TEXT_DARK
    p_main.font.name = FONT_FAMILY
    p_main.space_before = Pt(8)
    
    p_sub = tf1.add_paragraph()
    p_sub.text = "System Architecture Blueprint, Sequence Topologies, and Strategic Roadmap"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = TEXT_BODY
    p_sub.font.name = FONT_FAMILY
    p_sub.space_before = Pt(8)

    p_meta = tf1.add_paragraph()
    p_meta.text = f"Target Repository: {target_repo}   |   Standard: Tier-1 Production Standard   |   Verified Audit"
    p_meta.font.size = Pt(11.5)
    p_meta.font.color.rgb = TEXT_MUTED
    p_meta.font.name = FONT_FAMILY
    p_meta.space_before = Pt(28)

    add_slide_footer(slide1)

    # =========================================================
    # 2. Executive Mission & Problem Statement (Light Mode)
    # =========================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, prs)
    add_header(slide2, "Executive Mission & Problem Statement")
    add_card(slide2, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.3), 
             "Problem Statement & System Context", result.problem_statement, 
             font_size=13, title_color=ACCENT_BLUE)
    add_card(slide2, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3), 
             "Strategic Return on Investment (ROI)", 
             f"Executive Summary:\n{deck.executive_pitch}\n\nQuantified Business Value:\n{deck.business_value}", 
             font_size=13, title_color=ACCENT_GREEN)

    # =========================================================
    # 3. Strategy Matrix (4 Pillars with Distinct Light Accents)
    # =========================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, prs)
    add_header(slide3, "Strategic Architectural Deck Matrix")
    add_card(slide3, Inches(0.8), Inches(1.5), Inches(5.7), Inches(2.6), 
             "1. Executive Pitch", deck.executive_pitch, font_size=12, title_color=ACCENT_BLUE)
    add_card(slide3, Inches(6.8), Inches(1.5), Inches(5.7), Inches(2.6), 
             "2. Quantified Business Impact & ROI", deck.business_value, font_size=12, title_color=ACCENT_GREEN)
    add_card(slide3, Inches(0.8), Inches(4.3), Inches(5.7), Inches(2.6), 
             "3. Technical Challenges Solved", deck.technical_challenges, font_size=12, title_color=ACCENT_PURPLE)
    add_card(slide3, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.6), 
             "4. 12-Month Engineering Roadmap", deck.future_scope, font_size=12, title_color=ACCENT_AMBER)

    # =========================================================
    # 4. Quantitative Analytics & Capability Radar (Light Mode)
    # =========================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, prs)
    add_header(slide4, "Quantitative Analytics & Subsystem Density")
    metrics = result.analytics_metrics
    kpis = [
        ("SYSTEM COMPLEXITY", metrics.estimated_complexity.upper(), ACCENT_BLUE),
        ("AUDITED SUBSYSTEMS", f"{metrics.core_module_count} Modules", ACCENT_GREEN),
        ("SECURITY MODEL", metrics.security_posture.split(',')[0][:20], ACCENT_PURPLE),
        ("KNOWLEDGE NODES", f"{len(result.knowledge_items)} Nodes", ACCENT_AMBER)
    ]
    for i, (k_label, k_val, k_color) in enumerate(kpis):
        left_pos = Inches(0.8 + i * 2.95)
        k_card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.4), Inches(2.8), Inches(1.1))
        k_card.fill.solid()
        k_card.fill.fore_color.rgb = CARD_BG
        k_card.line.color.rgb = BORDER_COLOR
        k_card.line.width = Pt(1.2)

        # Top mini accent bar on KPI card
        top_k = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos + Inches(0.1), Inches(1.4), Inches(2.6), Inches(0.04))
        top_k.fill.solid()
        top_k.fill.fore_color.rgb = k_color
        top_k.line.color.rgb = k_color
        
        tb_k = slide4.shapes.add_textbox(left_pos + Inches(0.15), Inches(1.48), Inches(2.5), Inches(0.95))
        p1 = tb_k.text_frame.paragraphs[0]
        p1.text = k_label
        p1.font.size = Pt(9)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_MUTED
        p1.font.name = FONT_FAMILY
        
        p2 = tb_k.text_frame.add_paragraph()
        p2.text = k_val
        p2.font.size = Pt(20)
        p2.font.bold = True
        p2.font.color.rgb = k_color
        p2.font.name = FONT_FAMILY
        p2.space_before = Pt(2)

    try:
        c1_buf = generate_subsystem_breakdown_chart(result)
        slide4.shapes.add_picture(c1_buf, Inches(0.8), Inches(2.65), Inches(7.2), Inches(4.3))
    except Exception:
        pass

    try:
        c2_buf = generate_capability_radar_chart(result)
        slide4.shapes.add_picture(c2_buf, Inches(8.3), Inches(2.65), Inches(4.2), Inches(4.3))
    except Exception:
        pass

    # =========================================================
    # 5. System Architecture Blueprint (Light Mode Framing)
    # =========================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, prs)
    add_header(slide5, "System Architecture Blueprint")
    
    # White card frame for diagram
    diag_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.7), Inches(5.4))
    diag_card.fill.solid()
    diag_card.fill.fore_color.rgb = CARD_BG
    diag_card.line.color.rgb = BORDER_COLOR
    diag_card.line.width = Pt(1.2)

    arch_png_path = render_mermaid_to_png(result.architecture_mermaid_chart, f"{owner}_{repo}_arch")
    if arch_png_path and os.path.exists(arch_png_path):
        slide5.shapes.add_picture(arch_png_path, Inches(0.9), Inches(1.55), Inches(11.5), Inches(5.2))
    else:
        native_arch = generate_native_blueprint_chart(result.architecture_mermaid_chart, result.project_name)
        slide5.shapes.add_picture(native_arch, Inches(0.9), Inches(1.55), Inches(11.5), Inches(5.2))

    # =========================================================
    # 6. Primary Sequence Diagram (Light Mode Framing)
    # =========================================================
    if hg and getattr(hg, "execution_sequence_chart", None):
        slide6 = prs.slides.add_slide(blank_layout)
        set_slide_background(slide6, prs)
        add_header(slide6, "Primary Transaction & Workflow Sequence")
        
        diag_card2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.7), Inches(5.4))
        diag_card2.fill.solid()
        diag_card2.fill.fore_color.rgb = CARD_BG
        diag_card2.line.color.rgb = BORDER_COLOR
        diag_card2.line.width = Pt(1.2)

        seq_png_path = render_mermaid_to_png(hg.execution_sequence_chart, f"{owner}_{repo}_seq")
        if seq_png_path and os.path.exists(seq_png_path):
            slide6.shapes.add_picture(seq_png_path, Inches(0.9), Inches(1.55), Inches(11.5), Inches(5.2))
        else:
            native_seq = generate_native_sequence_chart(hg.execution_sequence_chart)
            slide6.shapes.add_picture(native_seq, Inches(0.9), Inches(1.55), Inches(11.5), Inches(5.2))

    # =========================================================
    # 7. Milestone Timeline (Light Mode Framing)
    # =========================================================
    if getattr(result, "milestone_timeline_mermaid_chart", None):
        slide7 = prs.slides.add_slide(blank_layout)
        set_slide_background(slide7, prs)
        add_header(slide7, "Architectural Evolution & Milestone Timeline")
        
        diag_card3 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.7), Inches(5.4))
        diag_card3.fill.solid()
        diag_card3.fill.fore_color.rgb = CARD_BG
        diag_card3.line.color.rgb = BORDER_COLOR
        diag_card3.line.width = Pt(1.2)

        time_png_path = render_mermaid_to_png(result.milestone_timeline_mermaid_chart, f"{owner}_{repo}_timeline")
        if time_png_path and os.path.exists(time_png_path):
            slide7.shapes.add_picture(time_png_path, Inches(0.9), Inches(1.55), Inches(11.5), Inches(5.2))
        else:
            add_card(slide7, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.3),
                     "System Milestones", result.milestone_timeline_mermaid_chart, font_size=12)

    # =========================================================
    # 8. Local Setup Prerequisites (Light Mode)
    # =========================================================
    if hg and hg.local_setup_prerequisites:
        slide8 = prs.slides.add_slide(blank_layout)
        set_slide_background(slide8, prs)
        add_header(slide8, "Engineering Handover: Setup & Execution")
        add_card(slide8, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.3),
                 "Local Setup Prerequisites & Configuration", hg.local_setup_prerequisites, font_size=12.5)

    # =========================================================
    # 9. Subsystem Directory Tour (Light Mode)
    # =========================================================
    if hg and hg.directory_tour:
        slide9 = prs.slides.add_slide(blank_layout)
        set_slide_background(slide9, prs)
        add_header(slide9, "Subsystem Directory Tour")
        dir_text = ""
        for d in hg.directory_tour[:10]:
            dir_text += f"• {d.folder_path}\n   {d.architectural_role}\n\n"
        add_card(slide9, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.3),
                 "Audited Functional Components", dir_text.strip(), font_size=12)

    # =========================================================
    # 10. Technical Debt & Verified Stack (Light Mode)
    # =========================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, prs)
    add_header(slide10, "Technical Debt & Verified Technology Stack")
    debt_content = ""
    if hg and hg.technical_debt_and_fragility:
        for t in hg.technical_debt_and_fragility:
            debt_content += f"• [PRIORITY WARNING] {t}\n\n"
    else:
        debt_content = "No high-severity technical fragility identified."
    add_card(slide10, Inches(0.8), Inches(1.5), Inches(6.5), Inches(5.3),
             "Technical Debt Priority Warnings", debt_content.strip(), font_size=12.5, 
             title_color=ACCENT_AMBER, bg_color=RGBColor(254, 252, 232), border_color=RGBColor(253, 224, 71))
    
    stack_content = f"Core Stack Components:\n{', '.join(result.tech_stack)}\n\n"
    stack_content += f"Architectural Patterns Detected:\n{', '.join(result.analytics_metrics.key_patterns_used)}\n\n"
    stack_content += f"Security Posture:\n{result.analytics_metrics.security_posture}"
    add_card(slide10, Inches(7.6), Inches(1.5), Inches(4.9), Inches(5.3),
             "Verified Technology Stack & Patterns", stack_content.strip(), font_size=12.5, 
             title_color=ACCENT_GREEN, bg_color=RGBColor(240, 253, 244), border_color=RGBColor(187, 247, 208))

    out_buf = BytesIO()
    prs.save(out_buf)
    out_buf.seek(0)
    return out_buf.getvalue()